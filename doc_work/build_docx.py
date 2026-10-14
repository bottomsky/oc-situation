from pathlib import Path
import re
import sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
source = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'finalized_content.md'
output = ROOT / '海上智能态势与筹划分析产品方案.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.54)
sec.left_margin = sec.right_margin = Cm(3.175)
sec.header_distance = sec.footer_distance = Cm(1.27)

def font_style(style, family, size, bold=False):
    style.font.name = family
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for name in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + name), family)
    for name in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        rf.attrib.pop(qn('w:' + name), None)

font_style(doc.styles['Normal'], '宋体', 12)
pf = doc.styles['Normal'].paragraph_format
pf.line_spacing = Pt(21)
pf.space_after = Pt(6)
pf.first_line_indent = Pt(24)
pf.widow_control = True
for name, size, before, after in [('Title', 22, 0, 18), ('Heading 1', 16, 15, 8), ('Heading 2', 13, 11, 6), ('Heading 3', 12, 8, 4)]:
    s = doc.styles[name]
    font_style(s, '黑体', size, False)
    p = s.paragraph_format
    p.first_line_indent = Pt(0)
    p.space_before, p.space_after = Pt(before), Pt(after)
    p.line_spacing = 1.25
    p.keep_with_next = True
    p.keep_together = True
    p.widow_control = True
font_style(doc.styles['List Paragraph'], '宋体', 12)
lp = doc.styles['List Paragraph'].paragraph_format
lp.first_line_indent = Pt(-12)
lp.left_indent = Pt(12)
lp.space_after = Pt(5)
lp.line_spacing = Pt(21)

def add_runs(p, text):
    for i, part in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        if part:
            p.add_run(part).bold = bool(i % 2)

def table(rows):
    count = len(rows[0])
    t = doc.add_table(rows=1, cols=count)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = ([2.4, 3.6, 4.35, 4.30] if count == 4 else [3.2, 5.75, 5.70] if count == 3 else [3.7, 10.95] if count == 2 else [14.65 / count] * count)
    for c, width in zip(t.columns, widths): c.width = Cm(width)
    pr = t._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for tag in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + tag)
        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: e.set(qn('w:' + key), val)
        b.append(e)
    pr.append(b)
    for ri, values in enumerate(rows):
        row = t.rows[0] if ri == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        for ci, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            mar = OxmlElement('w:tcMar')
            for tag, val in [('top', '90'), ('left', '110'), ('bottom', '90'), ('right', '110')]:
                e = OxmlElement('w:' + tag); e.set(qn('w:w'), val); e.set(qn('w:type'), 'dxa'); mar.append(e)
            cp.append(mar)
            if ri == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'F0F0F0'); cp.append(sh)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(17)
            p.paragraph_format.keep_with_next = ri == 0
            add_runs(p, text)
            for run in p.runs:
                run.font.size = Pt(10.5)
                if ri == 0: run.bold = True
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.line_spacing = Pt(3)
    tail.paragraph_format.first_line_indent = Pt(0)
    tail.add_run().font.size = Pt(3)

lines = source.read_text(encoding='utf-8-sig').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1; continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            values = [v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', v.replace(' ', '')) for v in values): rows.append(values)
            i += 1
        table(rows); continue
    h = re.match(r'^(#{1,4})\s+(.*)', line)
    if h:
        depth, content = len(h[1]), h[2]
        p = doc.add_paragraph(style='Title' if depth == 1 else 'Heading ' + str(depth - 1))
        add_runs(p, content)
        if depth == 1: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif re.match(r'^[-*]\s+', line):
        p = doc.add_paragraph(style='List Paragraph'); add_runs(p, '• ' + re.sub(r'^[-*]\s+', '', line))
    elif line == '---':
        pass
    else:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; add_runs(p, line)
    i += 1

# Simple page numbers provide navigation without decorative page furniture.
p = sec.footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Pt(0)
run = p.add_run()
run.font.size = Pt(9)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
r = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '1'; r.append(text); fld.append(r); p._p.append(fld)
doc.core_properties.title = '海上智能态势与筹划分析产品方案'
doc.core_properties.subject = '产品定位与核心能力'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.save(output)
print(output)
print('Paragraphs:', len(doc.paragraphs), 'Tables:', len(doc.tables), 'Characters:', len(source.read_text(encoding='utf-8-sig')))
